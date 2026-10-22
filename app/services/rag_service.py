"""RAG 服务：文档解析、中文切分、向量库维护与多轮问答链。

设计要点：
1. 单例管理，向量库与模型随首次使用懒加载，避免启动即加载 torch。
2. 向量 metadata 携带 document_id / filename，支持按文档精确删除。
3. 问答链使用 LangChain 1.x 的 LCEL 组合链（兼容层 langchain_classic）。
"""

from __future__ import annotations

import logging
import os
import threading
from pathlib import Path

from langchain_classic.chains import create_history_aware_retriever, create_retrieval_chain
from langchain_classic.chains.combine_documents import create_stuff_documents_chain
from langchain_core.documents import Document as LCDocument
from langchain_core.messages import AIMessage, BaseMessage, HumanMessage
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder

from app.config import CHUNK_SEPARATORS, get_settings

logger = logging.getLogger(__name__)

COLLECTION_NAME = "customer_kb"

# 未检索到依据时的固定话术（防幻觉核心）
FALLBACK_ANSWER = "您好，该问题暂未在知识库中找到相关内容，建议您联系人工客服处理，谢谢。"

SYSTEM_PROMPT = """你是「企业客服助手」，只依据下面提供的知识库片段回答客户问题。

规则：
1. 只能使用知识片段中的信息回答，禁止编造、推测或补充片段之外的内容。
2. 如果片段中没有能够回答该问题的依据，必须原样回复：{fallback}
3. 使用简体中文回答，语气专业、简洁，可在必要时分点说明。
4. 不要在回答中输出「参考来源」，系统会自动附加。

知识片段：
{context}
""".replace("{fallback}", FALLBACK_ANSWER)

CONDENSE_SYSTEM_PROMPT = """给定对话历史和用户的后续问题，请把后续问题改写成一个不依赖上下文、可独立检索的完整问题。

要求：只输出改写后的问句本身，不要解释、不要加引号；如果后续问题本身已经完整，则原样返回。"""


class RagService:
    """知识库 RAG 单例服务。"""

    _instance: "RagService | None" = None
    _instance_lock = threading.Lock()

    def __init__(self) -> None:
        self.settings = get_settings()
        self._lock = threading.Lock()
        self._client = None  # chromadb.PersistentClient，仅用于计数与集合管理
        self._store = None  # langchain_chroma.Chroma
        self._embeddings = None
        self._llm = None
        self._qa_chain = None

    # ------------------------------------------------------------------
    # 单例
    # ------------------------------------------------------------------
    @classmethod
    def instance(cls) -> "RagService":
        """获取全局单例（线程安全）。"""
        if cls._instance is None:
            with cls._instance_lock:
                if cls._instance is None:
                    cls._instance = cls()
        return cls._instance

    @classmethod
    def reset_instance(cls) -> None:
        """重置单例，主要供测试使用。"""
        with cls._instance_lock:
            cls._instance = None

    # ------------------------------------------------------------------
    # 懒加载组件
    # ------------------------------------------------------------------
    def _ensure_client(self):
        """初始化 Chroma 持久化客户端（不加载 embedding 模型）。"""
        if self._client is not None:
            return self._client
        import chromadb

        path = Path(self.settings.vector_store_dir)
        path.mkdir(parents=True, exist_ok=True)
        self._client = chromadb.PersistentClient(path=str(path))
        return self._client

    def _ensure_embeddings(self):
        """初始化向量化模型：默认本地 bge，可选通义 text-embedding-v3。"""
        if self._embeddings is not None:
            return self._embeddings

        provider = (self.settings.embedding_provider or "local").lower()
        if provider in {"dashscope", "tongyi", "qwen"}:
            if not self.settings.dashscope_api_key:
                raise RuntimeError(
                    "已选择通义向量模型，但未配置 DASHSCOPE_API_KEY。"
                    "请在 .env 中补充密钥，或将 EMBEDDING_PROVIDER 改回 local。"
                )
            from langchain_openai import OpenAIEmbeddings

            self._embeddings = OpenAIEmbeddings(
                model="text-embedding-v3",
                api_key=self.settings.dashscope_api_key,
                base_url=self.settings.dashscope_base_url,
                check_embedding_ctx_length=False,
            )
        else:
            # 未配置镜像时使用国内镜像，模型已预下载到本地缓存，可离线加载
            os.environ.setdefault("HF_ENDPOINT", "https://hf-mirror.com")
            from langchain_huggingface import HuggingFaceEmbeddings

            self._embeddings = HuggingFaceEmbeddings(
                model_name=self.settings.embedding_model,
                encode_kwargs={"normalize_embeddings": True},
            )
        return self._embeddings

    def _ensure_store(self):
        """初始化（并复用同一个 client 的）LangChain Chroma 实例。"""
        if self._store is not None:
            return self._store
        with self._lock:
            if self._store is None:
                from langchain_chroma import Chroma

                client = self._ensure_client()
                self._store = Chroma(
                    client=client,
                    collection_name=COLLECTION_NAME,
                    embedding_function=self._ensure_embeddings(),
                )
        return self._store

    def _ensure_llm(self):
        """初始化 DeepSeek 对话模型（OpenAI 兼容协议）。"""
        if self._llm is not None:
            return self._llm
        if not self.settings.llm_configured:
            raise RuntimeError(
                "未配置大模型密钥，请在 .env 中设置 DEEPSEEK_API_KEY 后重试。"
            )
        from langchain_openai import ChatOpenAI

        self._llm = ChatOpenAI(
            model=self.settings.deepseek_model,
            api_key=self.settings.deepseek_api_key,
            base_url=self.settings.deepseek_base_url,
            temperature=0.2,
            timeout=60,
            max_retries=1,
        )
        return self._llm

    # ------------------------------------------------------------------
    # 向量库读写
    # ------------------------------------------------------------------
    def collection_count(self) -> int:
        """当前集合中的向量数量（不加载 embedding 模型，可安全用于空库判断）。"""
        try:
            collection = self._ensure_client().get_or_create_collection(COLLECTION_NAME)
            return int(collection.count())
        except Exception as exc:  # 向量库损坏时应降级为「空库」，而不是让问答 500
            logger.warning("读取向量库数量失败：%s", exc)
            return 0

    # ------------------------------------------------------------------
    # 文档解析
    # ------------------------------------------------------------------
    def _load_file(self, local_path: Path, document_id: int, filename: str) -> list[LCDocument]:
        """按扩展名解析文件为 LangChain 文档列表。"""
        suffix = local_path.suffix.lower()
        base_meta = {"document_id": int(document_id), "filename": filename}
        docs: list[LCDocument] = []

        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(local_path))
            for index, page in enumerate(reader.pages, start=1):
                text = (page.extract_text() or "").strip()
                if text:
                    docs.append(
                        LCDocument(page_content=text, metadata={**base_meta, "page": index})
                    )
        elif suffix == ".docx":
            import docx

            document = docx.Document(str(local_path))
            paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            # 按段落合并成约 2000 字的块，避免段落过碎破坏语义；仍记录起始段落序号
            buffer: list[str] = []
            buffer_len = 0
            start_index = 1
            for index, para in enumerate(paragraphs, start=1):
                if not buffer:
                    start_index = index
                buffer.append(para)
                buffer_len += len(para)
                if buffer_len >= 2000:
                    docs.append(
                        LCDocument(
                            page_content="\n".join(buffer),
                            metadata={**base_meta, "page": start_index},
                        )
                    )
                    buffer, buffer_len = [], 0
            if buffer:
                docs.append(
                    LCDocument(
                        page_content="\n".join(buffer),
                        metadata={**base_meta, "page": start_index},
                    )
                )
        elif suffix in {".txt", ".md"}:
            text = local_path.read_text(encoding="utf-8", errors="ignore").strip()
            if text:
                docs.append(LCDocument(page_content=text, metadata={**base_meta, "page": 1}))
        else:
            raise ValueError(f"不支持的文件类型：{suffix}")

        if not docs:
            raise ValueError(
                f"文件「{filename}」未解析到文本内容，可能是扫描版 PDF 或空文档。"
            )
        return docs

    def _splitter(self):
        """中文友好的递归切分器，参数来自配置。"""
        from langchain_text_splitters import RecursiveCharacterTextSplitter

        return RecursiveCharacterTextSplitter(
            chunk_size=self.settings.chunk_size,
            chunk_overlap=self.settings.chunk_overlap,
            separators=CHUNK_SEPARATORS,
            length_function=len,
        )

    def add_document(self, document_id: int, filename: str, local_path: str | Path) -> int:
        """解析并索引单份文档，返回写入的文本块数量。"""
        path = Path(local_path)
        if not path.exists():
            raise ValueError(f"文件不存在：{path.name}")

        docs = self._load_file(path, document_id, filename)
        chunks = [c for c in self._splitter().split_documents(docs) if c.page_content.strip()]
        if not chunks:
            raise ValueError(f"文件「{filename}」切分后没有有效文本块。")

        # 统一补齐 metadata，保证后续可按 document_id 精确删除
        for chunk in chunks:
            chunk.metadata["document_id"] = int(document_id)
            chunk.metadata["filename"] = filename

        self._ensure_store().add_documents(chunks)
        logger.info("文档 %s（%s）索引完成，共 %s 个文本块", document_id, filename, len(chunks))
        return len(chunks)

    def delete_document(self, document_id: int) -> int:
        """删除某个文档的全部向量，返回删除数量。"""
        store = self._ensure_store()
        try:
            data = store.get(where={"document_id": int(document_id)})
        except Exception as exc:
            raise RuntimeError(f"查询向量失败：{exc}") from exc

        ids = list(data.get("ids") or [])
        if ids:
            store.delete(ids=ids)
        logger.info("文档 %s 向量删除完成，共 %s 条", document_id, len(ids))
        return len(ids)

    def rebuild_all(self, uploads_dir: str | Path) -> dict:
        """清空向量库后按磁盘文件全量重建索引。

        返回 {"total", "success", "failed", "items": [{document_id, chunk_count, error}]}
        """
        uploads = Path(uploads_dir)
        files = sorted([p for p in uploads.glob("*") if p.is_file() and p.suffix.lower() in {".pdf", ".docx", ".txt", ".md"}])

        with self._lock:
            if self._store is not None:
                try:
                    self._store.delete_collection()
                except Exception as exc:
                    logger.warning("清空集合失败，将直接重建：%s", exc)
            self._store = None
            try:
                self._ensure_client().delete_collection(COLLECTION_NAME)
            except Exception:
                # 集合本就不存在时忽略
                pass

        items: list[dict] = []
        for file in files:
            try:
                document_id = int(file.name.split("_", 1)[0])
            except (ValueError, IndexError):
                items.append(
                    {"document_id": None, "chunk_count": 0, "error": f"文件名缺少文档 ID 前缀：{file.name}"}
                )
                continue
            try:
                count = self.add_document(document_id, file.name.split("_", 1)[-1], file)
                items.append({"document_id": document_id, "chunk_count": count, "error": None})
            except Exception as exc:  # 单个文档失败不影响其他文档
                logger.warning("重建文档 %s 失败：%s", document_id, exc)
                items.append({"document_id": document_id, "chunk_count": 0, "error": str(exc)})

        success = sum(1 for item in items if item["error"] is None)
        return {
            "total": len(items),
            "success": success,
            "failed": len(items) - success,
            "items": items,
        }

    # ------------------------------------------------------------------
    # 问答
    # ------------------------------------------------------------------
    def _build_chain(self):
        """构建「历史改写 → 检索 → 生成」的 LCEL 问答链。"""
        if self._qa_chain is not None:
            return self._qa_chain

        llm = self._ensure_llm()
        retriever = self._ensure_store().as_retriever(
            search_kwargs={"k": self.settings.top_k}
        )

        condense_prompt = ChatPromptTemplate.from_messages(
            [
                ("system", CONDENSE_SYSTEM_PROMPT),
                MessagesPlaceholder("chat_history"),
                ("human", "{input}"),
            ]
        )
        history_aware_retriever = create_history_aware_retriever(llm, retriever, condense_prompt)

        qa_prompt = ChatPromptTemplate.from_messages(
            [
                ("system", SYSTEM_PROMPT),
                MessagesPlaceholder("chat_history"),
                ("human", "{input}"),
            ]
        )
        question_answer_chain = create_stuff_documents_chain(llm, qa_prompt)
        self._qa_chain = create_retrieval_chain(history_aware_retriever, question_answer_chain)
        return self._qa_chain

    @staticmethod
    def _to_messages(history_messages: list[dict] | None) -> list[BaseMessage]:
        """把数据库中的历史消息转换为 LangChain 消息对象。"""
        messages: list[BaseMessage] = []
        for item in history_messages or []:
            content = (item.get("content") or "").strip()
            if not content:
                continue
            if item.get("role") == "assistant":
                messages.append(AIMessage(content=content))
            else:
                messages.append(HumanMessage(content=content))
        return messages

    def ask(self, question: str, history_messages: list[dict] | None = None) -> dict:
        """回答用户问题。

        :return: {"answer": 回答文本, "sources": [文件名...]}
        """
        question = (question or "").strip()
        if not question:
            raise ValueError("问题内容不能为空")

        # 空知识库直接返回固定话术，不触发模型调用
        if self.collection_count() == 0:
            return {"answer": FALLBACK_ANSWER, "sources": []}

        chain = self._build_chain()
        try:
            result = chain.invoke(
                {"input": question, "chat_history": self._to_messages(history_messages)}
            )
        except Exception as exc:
            logger.exception("问答链执行失败")
            raise RuntimeError(f"智能客服暂时不可用，请稍后重试（{type(exc).__name__}）") from exc

        answer = (result.get("answer") or "").strip()
        if not answer:
            answer = FALLBACK_ANSWER

        sources: list[str] = []
        for doc in result.get("context") or []:
            name = (doc.metadata or {}).get("filename")
            if name and name not in sources:
                sources.append(name)

        if sources and "参考来源" not in answer:
            answer = f"{answer}\n\n参考来源：{'、'.join(sources)}"

        return {"answer": answer, "sources": sources}
